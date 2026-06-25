from flask import (Flask, render_template, request, redirect, url_for,
                   session, jsonify, Response, flash, abort, send_file)
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import generate_password_hash, check_password_hash
from functools import wraps
from datetime import datetime
import json, io, os, secrets, re
import openpyxl

# ─── App Setup ───────────────────────────────────────────────────────────────

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
INSTANCE_DIR = os.path.join(BASE_DIR, 'instance')
ABLAGE_DIR = os.path.join(INSTANCE_DIR, 'ablage')
TEMPLATE_XLSX = os.path.join(BASE_DIR, 'excel_template', 'EBD_Vorlage_leer.xlsx')
os.makedirs(INSTANCE_DIR, exist_ok=True)
os.makedirs(ABLAGE_DIR, exist_ok=True)


def _get_or_create_secret_key():
    key_file = os.path.join(INSTANCE_DIR, '.secret_key')
    if os.path.exists(key_file):
        with open(key_file) as f:
            return f.read().strip()
    key = secrets.token_hex(32)
    with open(key_file, 'w') as f:
        f.write(key)
    return key


app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', _get_or_create_secret_key())
app.config['SQLALCHEMY_DATABASE_URI'] = os.environ.get(
    'DATABASE_URL', f'sqlite:///{os.path.join(INSTANCE_DIR, "ebd.db")}')
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db = SQLAlchemy(app)

LISTE_INTERN = ['Rep. Leihgerät', 'Rep. Maschine BH', 'Ersatz Leihgerät', 'Ersatz Maschiene BH',
                'Muster', 'Kulanz', 'Schulungen/Kurse', 'Eigenbedarf', 'Schärfservice']
LISTE_EXTERN = ['Garantie', 'ext. Weiterverrechnung', 'Kulanz', 'Schärfservice']
LISTE_ABTEILUNG = ['GL & WWV', 'Wareneingang', 'Kasse / Warenausgang', 'SiFa', 'Parking / Shuttlebus',
                    'Eisenwaren /Nautic', 'Werkzeuge /Maschinen', 'Elektro', 'Sanitär', 'Fliesen',
                    'Baustoffe', 'Holz', 'Farben / Innendeko / Ambiente', 'Stadtgarten', 'Leihservice',
                    'Drive-In', 'Eigenbedarf Spezialfälle', 'Lehrlinge', 'Kundenzufuhr', 'NL', 'WK',
                    'Kassen & Systeme + Empfang', 'Einkauf', 'IT', 'FiCo', 'PA']

MWST_SATZ = 8.1

# ─── Models ──────────────────────────────────────────────────────────────────

class AdminUser(db.Model):
    __tablename__ = 'admin_users'
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password_hash = db.Column(db.String(256), nullable=False)

    def set_password(self, pw):
        self.password_hash = generate_password_hash(pw)

    def check_password(self, pw):
        return check_password_hash(self.password_hash, pw)


class EBDOrder(db.Model):
    __tablename__ = 'ebd_orders'
    id = db.Column(db.Integer, primary_key=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    bestell_nr = db.Column(db.String(50), nullable=False)
    firma = db.Column(db.String(200), default='')
    adresse = db.Column(db.Text, default='')
    verwendungszweck = db.Column(db.String(500), default='')
    art = db.Column(db.String(20), default='Intern')
    grund = db.Column(db.String(200), default='')
    besteller = db.Column(db.String(200), default='')
    abteilung = db.Column(db.String(200), default='')
    gl = db.Column(db.String(200), default='')
    mwst_satz = db.Column(db.Float, default=MWST_SATZ)
    items_json = db.Column(db.Text, default='[]')
    total_netto = db.Column(db.Float, default=0)
    total_mwst = db.Column(db.Float, default=0)
    total_brutto = db.Column(db.Float, default=0)
    excel_filename = db.Column(db.String(300), default='')

    @property
    def items(self):
        return json.loads(self.items_json or '[]')

    @items.setter
    def items(self, value):
        self.items_json = json.dumps(value, ensure_ascii=False)


# ─── Auth ────────────────────────────────────────────────────────────────────

def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'admin_id' not in session:
            return redirect(url_for('admin_login'))
        return f(*args, **kwargs)
    return decorated


def _slugify(value):
    value = re.sub(r'[^A-Za-z0-9_-]+', '-', value.strip())
    return value.strip('-') or 'firma'


def _next_bestell_nr():
    last = EBDOrder.query.order_by(EBDOrder.id.desc()).first()
    if last and last.bestell_nr.isdigit():
        return str(int(last.bestell_nr) + 1)
    return '731'


# ─── Excel generation ────────────────────────────────────────────────────────

def _generate_excel(order):
    wb = openpyxl.load_workbook(TEMPLATE_XLSX)
    ws = wb['EBD-Formular']

    address_lines = [l for l in (order.adresse or '').splitlines() if l.strip()][:4]
    for i, line in enumerate(address_lines):
        ws.cell(row=4 + i, column=6, value=line)  # F4..F7

    ws['E11'] = f'EBD-Bestell-Nr.: {order.bestell_nr}'
    ws['C17'] = order.verwendungszweck

    if order.art == 'Intern':
        ws['B18'] = order.grund
    else:
        ws['E18'] = order.grund

    items = order.items
    start_row = 20
    for i, item in enumerate(items[:5]):
        r = start_row + i
        ws.cell(row=r, column=1, value=item.get('artikelnummer', ''))
        ws.cell(row=r, column=3, value=item.get('artikelname', ''))
        ws.cell(row=r, column=6, value=item.get('menge', 0))
        ws.cell(row=r, column=7, value=item.get('einzelpreis', 0))
        ws.cell(row=r, column=8, value=f'=G{r}*F{r}')

    ws['H25'] = order.total_mwst
    ws['H26'] = order.total_brutto

    ws['B39'] = order.besteller
    ws['E39'] = order.abteilung
    ws['E41'] = order.gl

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _save_to_ablage(order, buf):
    year_dir = os.path.join(ABLAGE_DIR, str(order.created_at.year))
    os.makedirs(year_dir, exist_ok=True)
    filename = f'EBD_{order.bestell_nr}_{_slugify(order.firma)}_{order.created_at.strftime("%Y%m%d")}.xlsx'
    path = os.path.join(year_dir, filename)
    with open(path, 'wb') as f:
        f.write(buf.getvalue())
    order.excel_filename = os.path.join(str(order.created_at.year), filename)


# ─── Public Routes ───────────────────────────────────────────────────────────

@app.route('/')
def index():
    next_nr = _next_bestell_nr()
    return render_template('form.html', next_nr=next_nr,
                           liste_intern=LISTE_INTERN, liste_extern=LISTE_EXTERN,
                           liste_abteilung=LISTE_ABTEILUNG, mwst_satz=MWST_SATZ,
                           today=datetime.now().strftime('%d.%m.%Y'))


@app.route('/submit', methods=['POST'])
def submit():
    firma = request.form.get('firma', '').strip()
    adresse = request.form.get('adresse', '').strip()
    bestell_nr = request.form.get('bestell_nr', '').strip() or _next_bestell_nr()
    verwendungszweck = request.form.get('verwendungszweck', '').strip()
    art = request.form.get('art', 'Intern')
    grund = request.form.get('grund', '').strip()
    besteller = request.form.get('besteller', '').strip()
    abteilung = request.form.get('abteilung', '').strip()
    gl = request.form.get('gl', '').strip()
    mwst_satz = float(request.form.get('mwst_satz') or MWST_SATZ)

    if not firma or not verwendungszweck:
        flash('Bitte Firma und Verwendungszweck ausfüllen.', 'error')
        return redirect(url_for('index'))

    artikelnummern = request.form.getlist('artikelnummer[]')
    artikelnamen = request.form.getlist('artikelname[]')
    mengen = request.form.getlist('menge[]')
    preise = request.form.getlist('einzelpreis[]')

    items = []
    total_netto = 0.0
    for i in range(len(artikelnamen)):
        name = artikelnamen[i].strip()
        if not name:
            continue
        try:
            menge = float(mengen[i] or 0)
            preis = float(preise[i] or 0)
        except (ValueError, IndexError):
            menge, preis = 0, 0
        gesamt = menge * preis
        total_netto += gesamt
        items.append({'artikelnummer': artikelnummern[i] if i < len(artikelnummern) else '',
                      'artikelname': name, 'menge': menge, 'einzelpreis': preis, 'gesamtpreis': gesamt})

    if not items:
        flash('Bitte mindestens eine Artikelposition erfassen.', 'error')
        return redirect(url_for('index'))

    total_mwst = round(total_netto * mwst_satz / 100, 2)
    total_brutto = round(total_netto + total_mwst, 2)

    order = EBDOrder(bestell_nr=bestell_nr, firma=firma, adresse=adresse,
                     verwendungszweck=verwendungszweck, art=art, grund=grund,
                     besteller=besteller, abteilung=abteilung, gl=gl, mwst_satz=mwst_satz,
                     total_netto=round(total_netto, 2), total_mwst=total_mwst, total_brutto=total_brutto)
    order.items = items
    db.session.add(order)
    db.session.commit()

    buf = _generate_excel(order)
    _save_to_ablage(order, io.BytesIO(buf.getvalue()))
    db.session.commit()

    return redirect(url_for('success', oid=order.id))


@app.route('/success/<int:oid>')
def success(oid):
    order = EBDOrder.query.get_or_404(oid)
    return render_template('success.html', order=order)


@app.route('/download/<int:oid>')
def download(oid):
    order = EBDOrder.query.get_or_404(oid)
    buf = _generate_excel(order)
    filename = f'EBD_{order.bestell_nr}_{_slugify(order.firma)}.xlsx'
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ─── Admin Routes (Ablage / Archiv) ──────────────────────────────────────────

@app.route('/admin')
def admin_index():
    if 'admin_id' in session:
        return redirect(url_for('admin_dashboard'))
    return redirect(url_for('admin_login'))


@app.route('/admin/login', methods=['GET', 'POST'])
def admin_login():
    if 'admin_id' in session:
        return redirect(url_for('admin_dashboard'))
    error = None
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        user = AdminUser.query.filter_by(username=username).first()
        if user and user.check_password(password):
            session.permanent = True
            session['admin_id'] = user.id
            session['admin_username'] = user.username
            return redirect(url_for('admin_dashboard'))
        error = 'Ungültige Anmeldedaten.'
    return render_template('admin/login.html', error=error)


@app.route('/admin/logout')
def admin_logout():
    session.clear()
    return redirect(url_for('admin_login'))


@app.route('/admin/dashboard')
@login_required
def admin_dashboard():
    page = request.args.get('page', 1, type=int)
    search = request.args.get('search', '').strip()
    query = EBDOrder.query.order_by(EBDOrder.created_at.desc())
    if search:
        like = f'%{search}%'
        query = query.filter(db.or_(EBDOrder.firma.ilike(like),
                                    EBDOrder.bestell_nr.ilike(like),
                                    EBDOrder.verwendungszweck.ilike(like)))
    orders = query.paginate(page=page, per_page=25, error_out=False)
    total = EBDOrder.query.count()
    total_brutto_sum = db.session.query(db.func.sum(EBDOrder.total_brutto)).scalar() or 0
    return render_template('admin/dashboard.html', orders=orders, total=total,
                           total_brutto_sum=total_brutto_sum, search=search)


@app.route('/admin/ablage/<int:oid>/download')
@login_required
def admin_download(oid):
    order = EBDOrder.query.get_or_404(oid)
    if order.excel_filename:
        path = os.path.join(ABLAGE_DIR, order.excel_filename)
        if os.path.exists(path):
            return send_file(path, as_attachment=True,
                             download_name=os.path.basename(path),
                             mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    buf = _generate_excel(order)
    filename = f'EBD_{order.bestell_nr}_{_slugify(order.firma)}.xlsx'
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/admin/ablage/<int:oid>/delete', methods=['POST'])
@login_required
def admin_delete_order(oid):
    order = EBDOrder.query.get_or_404(oid)
    if order.excel_filename:
        path = os.path.join(ABLAGE_DIR, order.excel_filename)
        if os.path.exists(path):
            os.remove(path)
    db.session.delete(order)
    db.session.commit()
    flash('Bestellung aus der Ablage gelöscht.', 'success')
    return redirect(url_for('admin_dashboard'))


@app.route('/admin/change-password', methods=['POST'])
@login_required
def admin_change_password():
    user = AdminUser.query.get(session['admin_id'])
    current = request.form.get('current_password', '')
    new_pw = request.form.get('new_password', '')
    confirm = request.form.get('confirm_password', '')
    if not user.check_password(current):
        flash('Aktuelles Passwort ist falsch.', 'error')
    elif new_pw != confirm:
        flash('Neue Passwörter stimmen nicht überein.', 'error')
    elif len(new_pw) < 6:
        flash('Passwort muss mindestens 6 Zeichen lang sein.', 'error')
    else:
        user.set_password(new_pw)
        db.session.commit()
        flash('Passwort erfolgreich geändert!', 'success')
    return redirect(url_for('admin_dashboard'))


# ─── Init Helpers ────────────────────────────────────────────────────────────

def init_db():
    db.create_all()
    if not AdminUser.query.first():
        admin_user = os.environ.get('ADMIN_USERNAME', 'admin')
        admin_pass = os.environ.get('ADMIN_PASSWORD', 'admin123')
        user = AdminUser(username=admin_user)
        user.set_password(admin_pass)
        db.session.add(user)
        print(f'[INIT] Admin-Benutzer erstellt: {admin_user} / {admin_pass}')
        print('[INIT] Bitte Passwort nach dem ersten Login ändern!')
    db.session.commit()


if __name__ == '__main__':
    with app.app_context():
        init_db()
    port = int(os.environ.get('PORT', 5000))
    host = os.environ.get('HOST', '0.0.0.0')
    debug = os.environ.get('FLASK_DEBUG', 'false').lower() == 'true'
    app.run(debug=debug, host=host, port=port)
